import argparse
from pathlib import Path
from openpyxl import load_workbook
from urllib.parse import urlparse
from utils import normalize_cell, normalize_email, is_valid_email, write_csv

OUTPUT_FIELDS = [
    "email",
    "first_name",
    "last_name",
    "agency_name",
    "agency_label",
    "website",
    "website_safe",
    "domain",
    "city",
    "province",
    "country",
    "phone",
    "language",
    "source",
]


def iter_rows_xlsx(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            continue
        header = [normalize_cell(h) for h in header]
        col_index = {name: i for i, name in enumerate(header) if name}
        for row in rows:
            yield sheet, col_index, row


SOCIAL_DOMAINS = {
    "linkedin.com",
    "facebook.com",
    "instagram.com",
    "twitter.com",
    "x.com",
}


def domain_from_url(url):
    if not url:
        return ""
    try:
        parsed = urlparse(url)
        host = parsed.netloc or parsed.path
        host = host.replace("www.", "")
        return host
    except Exception:
        return ""


def agency_label(agency_name, website):
    name = (agency_name or "").strip()
    if name and not name.lower().startswith("http") and "linkedin.com" not in name.lower():
        return name
    domain = domain_from_url(website)
    if domain and domain not in SOCIAL_DOMAINS:
        return domain
    if name and not name.lower().startswith("http"):
        return name
    return ""


def website_safe(website):
    domain = domain_from_url(website)
    if not domain or domain in SOCIAL_DOMAINS:
        return ""
    return website


def build_from_leads_hispano(path, include_countries=None):
    rows_out = []
    for sheet, cols, row in iter_rows_xlsx(path):
        def get(name):
            if name not in cols:
                return ""
            return normalize_cell(row[cols[name]])
        email = normalize_email(get("Email"))
        if not is_valid_email(email):
            continue
        country = get("Country")
        if include_countries:
            if not country or country not in include_countries:
                continue
        agency_name = get("Agency")
        website = get("Website")
        rows_out.append({
            "email": email,
            "first_name": get("Name"),
            "last_name": get("Last Name"),
            "agency_name": agency_name,
            "agency_label": agency_label(agency_name, website),
            "website": website,
            "website_safe": website_safe(website),
            "domain": domain_from_url(website),
            "city": get("City"),
            "province": get("State/Province"),
            "country": country,
            "phone": get("Phone Number"),
            "language": "es",
            "source": get("Source Sheet") or sheet,
        })
    return rows_out


def split_country_from_province(province):
    if not province:
        return "", ""
    if "," in province:
        parts = [p.strip() for p in province.split(",")]
        if len(parts) >= 2:
            return ", ".join(parts[:-1]), parts[-1]
    return province, ""


def build_from_agencies_all_provinces(path):
    rows_out = []
    for sheet, cols, row in iter_rows_xlsx(path):
        def get(name):
            if name not in cols:
                return ""
            return normalize_cell(row[cols[name]])
        email = normalize_email(get("email"))
        if not is_valid_email(email):
            continue
        province_raw = get("province")
        province, country = split_country_from_province(province_raw)
        agency_name = get("name")
        website = get("website")
        rows_out.append({
            "email": email,
            "first_name": "",
            "last_name": "",
            "agency_name": agency_name,
            "agency_label": agency_label(agency_name, website),
            "website": website,
            "website_safe": website_safe(website),
            "domain": domain_from_url(website),
            "city": get("city"),
            "province": province,
            "country": country,
            "phone": get("phone"),
            "language": "es",
            "source": sheet,
        })
    return rows_out


def dedupe_by_email(rows):
    seen = set()
    out = []
    for r in rows:
        email = r["email"].lower()
        if email in seen:
            continue
        seen.add(email)
        out.append(r)
    return out


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--leads-hispano", default="Leads_Hispano_Hablantes.xlsx")
    parser.add_argument("--agencies", default="")
    parser.add_argument(
        "--countries",
        default="Spain,Mexico,Argentina,Colombia,Chile,Peru,Ecuador,Uruguay,Paraguay,Bolivia,Guatemala,Honduras,El Salvador,Nicaragua,Costa Rica,Panama,Puerto Rico,Dominican Republic,Venezuela,Cuba",
        help="Comma-separated list of countries to include (exact match).",
    )
    parser.add_argument("--out", default="gtm/output/leads_clean.csv")
    args = parser.parse_args()

    include_countries = [c.strip() for c in args.countries.split(",") if c.strip()]

    rows = []
    if Path(args.leads_hispano).exists():
        rows += build_from_leads_hispano(args.leads_hispano, include_countries)
    if args.agencies and Path(args.agencies).exists():
        rows += build_from_agencies_all_provinces(args.agencies)

    rows = dedupe_by_email(rows)

    write_csv(args.out, rows, OUTPUT_FIELDS)
    print(f"wrote {len(rows)} rows to {args.out}")


if __name__ == "__main__":
    main()
